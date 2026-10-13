from io import BytesIO
from django.test import TestCase
from django.contrib.auth.models import User,Permission
from openpyxl import load_workbook
from .models import Laboratorio,Equipo,Mantenimiento,Solicitud

class ReporteEstadisticasTests(TestCase):
    def setUp(self):
        self.user=User.objects.create_user('reportes')
        self.user.user_permissions.add(Permission.objects.get(content_type__app_label='lab',codename='view_mantenimiento'))
        self.client.force_login(self.user)
        self.lab=Laboratorio.objects.create(nombre='603')
        self.equipo=Equipo.objects.create(laboratorio=self.lab,descripcion='Equipo',placa='001')
        Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.user,tipo='MIC',fecha='2026-10-13',diagnostico='=SUM(1,2)')
        Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.user,tipo='MIPV',fecha='2026-09-01')
        Solicitud.objects.create(equipo=self.equipo,solicitante=self.user,nombre_solicitante='Docente',tipo='MIC',descripcion='Falla')
    def workbook(self,filters):
        response=self.client.get('/mantenimientos/estadisticas/excel/',filters)
        self.assertEqual(response.status_code,200)
        self.assertIn('attachment',response['Content-Disposition'])
        return load_workbook(BytesIO(b''.join(response.streaming_content)))
    def test_reporte_coincide_con_filtros_y_conserva_texto(self):
        filters={'laboratorio':self.lab.pk,'tipo':'MIC','desde':'2026-10-01','hasta':'2026-10-31'}
        response=self.client.get('/mantenimientos/estadisticas/',filters)
        self.assertContains(response,'Descargar reporte Excel')
        wb=self.workbook(filters)
        self.assertEqual(len(wb.sheetnames),5)
        self.assertEqual(wb['Resumen']['B3'].value,response.context['total'])
        self.assertEqual(wb['Por tipo']['B2'].value,1)
        self.assertEqual(wb['Por laboratorio']['A2'].value,'603')
        self.assertEqual(wb['Intervenciones'].max_row,2)
        self.assertEqual(wb['Intervenciones']['H2'].value,'=SUM(1,2)')
        self.assertEqual(wb['Intervenciones']['H2'].data_type,'s')
        self.assertEqual(wb['Solicitudes todos los períodos']['B2'].value,1)
        wb.close()
    def test_reporte_vacio_y_permisos(self):
        wb=self.workbook({'q':'Sin coincidencias'})
        self.assertEqual(wb['Resumen']['B3'].value,0)
        self.assertEqual(wb['Intervenciones'].max_row,1)
        wb.close()
        self.client.force_login(User.objects.create_user('sin_permiso'))
        self.assertEqual(self.client.get('/mantenimientos/estadisticas/excel/').status_code,403)
        self.client.logout()
        self.assertEqual(self.client.get('/mantenimientos/estadisticas/excel/').status_code,302)
