from unittest.mock import patch
from django.contrib.auth.models import User, Permission
from django.test import TestCase
from django.urls import reverse
from .models import Laboratorio, Equipo, Traslado, Auditoria
from .transfers import ESTADO_TRASLADO

class TrasladoTests(TestCase):
    def setUp(self):
        self.a=User.objects.create_user('tecnico_a')
        self.b=User.objects.create_user('tecnico_b')
        self.c=User.objects.create_user('tercero')
        permissions=Permission.objects.filter(content_type__app_label='lab',codename__in=['view_equipo','change_equipo'])
        for user in [self.a,self.b,self.c]: user.user_permissions.set(permissions)
        self.origen=Laboratorio.objects.create(nombre='Origen',responsable=self.b)
        self.destino=Laboratorio.objects.create(nombre='Destino',responsable=self.a)
        self.equipo=Equipo.objects.create(descripcion='Osciloscopio',serial='SERIAL-123',laboratorio=self.origen,estado='Operativo')
        self.url=reverse('pedir_traslado',args=[self.equipo.pk])
        self.payload={'destino':self.destino.pk,'justificacion':'Práctica de electrónica'}
        self.client.force_login(self.a)

    def pedir(self):
        self.assertEqual(self.client.post(self.url,self.payload).status_code,302)
        return Traslado.objects.get()

    def test_solicitud_alerta_y_tarjeta(self):
        t=self.pedir()
        self.equipo.refresh_from_db()
        self.assertEqual(self.equipo.estado,ESTADO_TRASLADO)
        self.assertEqual(self.equipo.laboratorio,self.origen)
        self.client.force_login(self.b)
        self.assertEqual(self.client.get(reverse('alertas_traslado')).json(),{'pendientes':1,'solicitudes':0})
        response=self.client.get(reverse('traslados'))
        for text in ['Origen','Destino','SERIAL-123','Práctica de electrónica','Aprobar','Rechazar']:
            self.assertContains(response,text)
        self.assertEqual(t.aprobador,self.b)

    def test_accesos_y_explicacion_sin_responsable(self):
        self.origen.responsable=None; self.origen.save()
        response=self.client.get(reverse('detalle',args=[self.equipo.pk]))
        self.assertContains(response,'Pedir traslado')
        self.assertContains(response,'Primero asigna un responsable')
        admin_user=User.objects.create_superuser('admin_ui',password='test')
        self.client.force_login(admin_user)
        self.assertContains(self.client.get(reverse('equipos')),'Asignar salones a t&#233;cnicos')
        self.assertContains(self.client.get(reverse('laboratorios')),'Usuarios y roles')

    def test_aprobar_y_no_repetir(self):
        t=self.pedir(); self.client.force_login(self.b)
        url=reverse('resolver_traslado',args=[t.pk])
        self.assertEqual(self.client.get(url).status_code,405)
        self.client.post(url,{'decision':'Aprobado'})
        self.equipo.refresh_from_db(); t.refresh_from_db()
        self.assertEqual(self.equipo.laboratorio,self.destino)
        self.assertEqual(self.equipo.estado,'Operativo')
        self.assertEqual(t.estado,'Aprobado')
        self.assertIsNotNone(t.resuelto)
        self.assertEqual(self.client.get(reverse('alertas_traslado')).json(),{'pendientes':0,'solicitudes':0})
        self.client.post(url,{'decision':'Rechazado'})
        t.refresh_from_db(); self.assertEqual(t.estado,'Aprobado')
        self.assertEqual(Auditoria.objects.count(),2)

    def test_rechazar_restaurar(self):
        t=self.pedir(); self.client.force_login(self.b)
        self.client.post(reverse('resolver_traslado',args=[t.pk]),{'decision':'Rechazado','respuesta':'Necesario aquí'})
        self.equipo.refresh_from_db(); t.refresh_from_db()
        self.assertEqual(self.equipo.laboratorio,self.origen)
        self.assertEqual(self.equipo.estado,'Operativo')
        self.assertEqual(t.respuesta,'Necesario aquí')

    def test_permisos_y_privacidad(self):
        self.assertEqual(self.client.get(reverse('editar_equipo',args=[self.equipo.pk])).status_code,403)
        t=self.pedir()
        for user in [self.a,self.c]:
            self.client.force_login(user)
            self.assertEqual(self.client.post(reverse('resolver_traslado',args=[t.pk]),{'decision':'Aprobado'}).status_code,403)
        self.assertNotContains(self.client.get(reverse('traslados')),'SERIAL-123')
        self.client.force_login(self.b)
        self.assertEqual(self.client.get(reverse('editar_equipo',args=[self.equipo.pk])).status_code,403)

    def test_destino_ajeno_vacio_y_duplicado(self):
        for payload in [{**self.payload,'destino':self.origen.pk},{**self.payload,'justificacion':'   '}]:
            self.assertEqual(self.client.post(self.url,payload).status_code,200)
            self.assertFalse(Traslado.objects.exists())
        self.pedir()
        self.assertContains(self.client.post(self.url,self.payload),'ya tiene un traslado pendiente')
        self.assertEqual(Traslado.objects.count(),1)

    def test_guardado_atomico(self):
        with patch('lab.views.audit',side_effect=RuntimeError('Error de auditoría')):
            with self.assertRaises(RuntimeError): self.client.post(self.url,self.payload)
        self.equipo.refresh_from_db()
        self.assertEqual(self.equipo.estado,'Operativo')
        self.assertFalse(Traslado.objects.exists())

    def test_respuesta_atomica_y_csrf(self):
        from django.test import Client
        t=self.pedir(); self.client.force_login(self.b)
        url=reverse('resolver_traslado',args=[t.pk])
        csrf_client=Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.b)
        self.assertEqual(csrf_client.post(url,{'decision':'Aprobado'}).status_code,403)
        with patch('lab.views.audit',side_effect=RuntimeError('Error de auditoría')):
            with self.assertRaises(RuntimeError): self.client.post(url,{'decision':'Aprobado'})
        t.refresh_from_db(); self.equipo.refresh_from_db()
        self.assertEqual(t.estado,'Pendiente')
        self.assertEqual(self.equipo.estado,ESTADO_TRASLADO)
        self.assertEqual(self.equipo.laboratorio,self.origen)

    def test_fotografia_hoja_y_admin_respetan_responsable(self):
        from django.contrib import admin
        from django.test import RequestFactory
        for name in ['editar_hoja','cambiar_foto']:
            self.assertEqual(self.client.get(reverse(name,args=[self.equipo.pk])).status_code,403)
        request=RequestFactory().get('/admin/'); request.user=self.a
        equipo_admin=admin.site._registry[Equipo]
        self.assertFalse(equipo_admin.has_change_permission(request,self.equipo))
        request.user=self.b
        self.assertTrue(equipo_admin.has_change_permission(request,self.equipo))
        self.pedir()
        self.assertFalse(equipo_admin.has_change_permission(request,self.equipo))

    def test_cambio_responsable_y_decision_invalida(self):
        t=self.pedir(); self.client.force_login(self.b)
        url=reverse('resolver_traslado',args=[t.pk])
        self.client.post(url,{'decision':'Inventada'})
        t.refresh_from_db(); self.assertEqual(t.estado,'Pendiente')
        self.destino.responsable=self.c; self.destino.save()
        self.client.post(url,{'decision':'Aprobado'})
        t.refresh_from_db(); self.assertEqual(t.estado,'Pendiente')
        self.client.post(url,{'decision':'Rechazado'})
        t.refresh_from_db(); self.assertEqual(t.estado,'Rechazado')

    def test_no_mover_desde_edicion(self):
        self.client.force_login(self.b)
        response=self.client.post(reverse('editar_equipo',args=[self.equipo.pk]),{'descripcion':'Osciloscopio','laboratorio':self.destino.pk,'version':1})
        self.assertEqual(response.status_code,302)
        self.equipo.refresh_from_db(); self.assertEqual(self.equipo.laboratorio,self.origen)

    def test_reporte_respeta_privacidad_y_texto(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.payload['justificacion']='=HYPERLINK("https://example.com")'
        self.pedir()
        response=self.client.get(reverse('exportar_traslados'))
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet['L2'].value,self.payload['justificacion'])
        self.assertEqual(sheet['L2'].data_type,'s')
        self.client.force_login(self.c)
        sheet=load_workbook(BytesIO(self.client.get(reverse('exportar_traslados')).content)).active
        self.assertEqual(sheet.max_row,1)

    def test_inicio_centraliza_laboratorios_propios(self):
        response=self.client.get(reverse('inicio'))
        self.assertContains(response,'Gestión de tus laboratorios')
        self.assertEqual(list(response.context['mis_laboratorios']),[self.destino])
        self.assertEqual(response.context['metrics']['Equipos'],0)
        self.assertContains(response,'dashboard-card')
