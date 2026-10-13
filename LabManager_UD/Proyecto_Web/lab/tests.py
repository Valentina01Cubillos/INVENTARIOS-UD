from datetime import date
from django.test import TestCase,override_settings
from django.contrib.auth.models import User,Permission
from .models import Equipo,Laboratorio,Solicitud,Mantenimiento
from unittest import skipUnless
from django.db import connection, close_old_connections
from django.test import TransactionTestCase
class FlujosTests(TestCase):
    def setUp(self):
        self.user=User.objects.create_user('solicitante',password='Password-test-935!')
        self.other=User.objects.create_user('otro',password='Password-test-936!')
        self.equipo=Equipo.objects.create(descripcion='Multímetro',laboratorio=Laboratorio.objects.create(nombre='603'),placa='001')
    def grant(self,*codes):
        self.user.user_permissions.add(*Permission.objects.filter(content_type__app_label='lab',codename__in=codes))
    def test_login_y_permisos(self):
        self.assertEqual(self.client.get('/equipos/').status_code,302)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/equipos/').status_code,403)

    def test_solicitante_no_se_reemplaza_por_cuenta_tecnica(self):
        sol=Solicitud.objects.create(solicitante=self.user,tecnico=self.user,tipo='MIC',descripcion='Revisar',nombre_solicitante=self.user.username)
        self.client.force_login(self.user)
        self.assertContains(self.client.get('/solicitudes/'),'Solicitante por confirmar')
        url=f'/solicitudes/{sol.pk}/solicitante/'
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(url,{'nombre_solicitante':'Docente real','version':sol.version}).status_code,403)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(url).status_code,405)
        self.client.post(url,{'nombre_solicitante':'Docente real','version':sol.version})
        sol.refresh_from_db()
        self.assertEqual(sol.solicitante_nombre,'Docente real')
        self.assertEqual(sol.tecnico,self.user)
        self.assertEqual(sol.solicitante,self.user)
        self.assertContains(self.client.get('/solicitudes/'),'Docente real')
        self.assertContains(self.client.get(f'/solicitudes/{sol.pk}/'),'Docente real')

    def test_diagnostico_obligatorio_observaciones_opcionales(self):
        self.client.force_login(self.user)
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.other,tecnico=self.user,estado='Asignada',tipo='MIC',descripcion='Falla')
        payload={'solicitud':sol.pk,'equipo':self.equipo.pk,'tipo':'MIC','fecha':'2026-10-13','diagnostico':'   '}
        response=self.client.post('/mantenimientos/nuevo/',payload)
        self.assertIn('diagnostico',response.context['form'].errors)
        self.assertNotIn('observaciones',response.context['form'].errors)
        self.assertFalse(Mantenimiento.objects.exists())
        payload['diagnostico']='Se encontró suciedad en los contactos.'
        self.assertEqual(self.client.post('/mantenimientos/nuevo/',payload).status_code,302)
        self.assertEqual(Mantenimiento.objects.get().observaciones,'')

    def test_eliminar_solicitud_confirmacion_y_permisos(self):
        from django.test import Client
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.user,tipo='MIC',descripcion='Prueba')
        url=f'/solicitudes/{sol.pk}/eliminar/'
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(url).status_code,403)
        self.assertEqual(self.client.post(url,{'version':sol.version,'confirmar':'si'}).status_code,403)
        self.client.force_login(self.user)
        self.assertContains(self.client.get(f'/solicitudes/{sol.pk}/'),'Eliminar solicitud')
        self.assertContains(self.client.get(url),'Confirmo que quiero eliminar')
        self.assertTrue(Solicitud.objects.filter(pk=sol.pk).exists())
        csrf=Client(enforce_csrf_checks=True);csrf.force_login(self.user)
        self.assertEqual(csrf.post(url,{'version':sol.version,'confirmar':'si'}).status_code,403)
        self.client.post(url,{'version':sol.version})
        self.client.post(url,{'version':0,'confirmar':'si'})
        self.assertTrue(Solicitud.objects.filter(pk=sol.pk).exists())
        self.assertRedirects(self.client.post(url,{'version':sol.version,'confirmar':'si'}),'/solicitudes/')
        self.assertFalse(Solicitud.objects.filter(pk=sol.pk).exists())

    def test_eliminar_conserva_mantenimiento_y_no_reutiliza_numero(self):
        from .models import Auditoria
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.user,tecnico=self.user,tipo='MIC',descripcion='Original',estado='Cerrada')
        numero=sol.codigo
        m=Mantenimiento.objects.create(equipo=self.equipo,solicitud=sol,tecnico=self.user,tipo='MIC',especificaciones='Trabajo conservado',datos_legacy={'original':'dato'})
        self.client.force_login(self.user)
        self.client.post(f'/solicitudes/{sol.pk}/eliminar/',{'version':sol.version,'confirmar':'si'})
        m.refresh_from_db()
        self.assertIsNone(m.solicitud_id)
        self.assertEqual(m.especificaciones,'Trabajo conservado')
        self.assertEqual(m.datos_legacy['codigo_solicitud'],numero)
        self.assertEqual(m.datos_legacy['original'],'dato')
        self.assertContains(self.client.get(f'/mantenimientos/{m.pk}/'),numero)
        self.assertTrue(Auditoria.objects.filter(accion='Eliminar solicitud').exists())
        nueva=Solicitud.objects.create(solicitante=self.user,tipo='MIC',descripcion='Nueva')
        self.assertGreater(int(nueva.codigo[2:]),int(numero[2:]))

    def test_liberar_mantenimiento_disponible_inmediatamente(self):
        from django.test import Client
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.other,tipo='MIC',descripcion='Falla',tecnico=self.user,estado='Asignada')
        url=f'/solicitudes/{sol.pk}/liberar/'
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(url).status_code,403)
        self.client.force_login(self.user)
        listado=self.client.get('/solicitudes/')
        self.assertContains(listado,'Continuar mantenimiento')
        self.assertNotContains(listado,'Cancelar mantenimiento')
        self.assertContains(self.client.get(f'/solicitudes/{sol.pk}/'),'Continuar mantenimiento')
        self.assertEqual(self.client.get(url).status_code,405)
        csrf_client=Client(enforce_csrf_checks=True);csrf_client.force_login(self.user)
        self.assertEqual(csrf_client.post(url).status_code,403)
        self.assertContains(self.client.get(f'/mantenimientos/nuevo/?solicitud={sol.pk}'),'Cancelar mantenimiento')
        self.assertRedirects(self.client.post(url),'/solicitudes/')
        sol.refresh_from_db()
        self.assertIsNone(sol.tecnico_id)
        self.assertEqual(sol.estado,'Pendiente')
        self.assertFalse(Mantenimiento.objects.filter(solicitud=sol).exists())
        self.client.post(f'/mantenimientos/nuevo/?solicitud={sol.pk}',{'especificaciones':'Formulario viejo','estado':'Operativo'})
        self.assertFalse(Mantenimiento.objects.filter(solicitud=sol).exists())
        self.client.force_login(self.other)
        self.assertContains(self.client.get('/solicitudes/'),'Realizar mantenimiento')
        self.client.post(f'/solicitudes/{sol.pk}/realizar/')
        sol.refresh_from_db();self.assertEqual(sol.tecnico,self.other)
        sol.estado='Cerrada';sol.save()
        self.client.post(url)
        sol.refresh_from_db();self.assertEqual(sol.estado,'Cerrada')

    def test_registro_directo_asigna_numero_automatico(self):
        Solicitud.objects.create(solicitante=self.other,tipo='MIC',descripcion='Histórica',codigo_legacy='S-164')
        self.client.force_login(self.user)
        response=self.client.get('/mantenimientos/nuevo/')
        self.assertContains(response,'se asignará automáticamente al guardar')
        self.assertNotIn('solicitud',response.context['form'].fields)
        self.assertNotIn('tecnico',response.context['form'].fields)
        for i in range(2):
            response=self.client.post('/mantenimientos/nuevo/',{'equipo':self.equipo.pk,'nombre_solicitante':'Docente solicitante','tipo':'MIC','descripcion':'Revisión del equipo'})
            self.assertEqual(response.status_code,302)
            solicitud=Solicitud.objects.latest('pk')
            self.assertEqual(solicitud.estado,'Pendiente')
            self.assertIsNone(solicitud.tecnico_id)
            self.assertEqual(solicitud.nombre_solicitante,'Docente solicitante')
            self.assertEqual(solicitud.solicitante,self.user)
            self.assertFalse(Mantenimiento.objects.exists())
            self.assertContains(self.client.get(response.url),solicitud.codigo)
        solicitudes=list(Solicitud.objects.exclude(codigo_legacy='S-164').order_by('pk'))
        self.assertEqual([s.codigo for s in solicitudes],['S-165','S-166'])
        self.assertEqual(solicitudes[1].pk,solicitudes[0].pk+1)
        self.assertNotEqual(solicitudes[0].codigo,solicitudes[1].codigo)
        self.client.force_login(self.other)
        self.assertContains(self.client.get('/solicitudes/'),'Realizar mantenimiento')
        self.client.post(f'/solicitudes/{solicitudes[0].pk}/realizar/')
        solicitudes[0].refresh_from_db()
        self.assertEqual(solicitudes[0].tecnico,self.other)
        self.assertEqual(solicitudes[0].estado,'Asignada')
        self.assertEqual(solicitudes[0].nombre_solicitante,'Docente solicitante')
    def test_inventario_por_laboratorio_y_busqueda(self):
        self.grant('view_equipo'); self.client.force_login(self.user)
        other=Laboratorio.objects.create(nombre='604')
        Equipo.objects.create(descripcion='Equipo ajeno',laboratorio=other)
        response=self.client.get('/equipos/')
        self.assertContains(response,f'?laboratorio={self.equipo.laboratorio_id}')
        self.assertNotContains(response,'Ver equipo</a>')
        response=self.client.get('/equipos/',{'laboratorio':self.equipo.laboratorio_id})
        self.assertContains(response,'Ver equipo</a>')
        self.assertNotContains(response,'Equipo ajeno')
        response=self.client.get('/equipos/',{'q':'Equipo ajeno'})
        self.assertEqual(list(response.context['page']),list(Equipo.objects.filter(laboratorio=other)))
        response=self.client.get('/equipos/',{'laboratorio':self.equipo.laboratorio_id,'q':'Equipo ajeno'})
        self.assertEqual(response.context['page'].paginator.count,0)
        self.assertEqual(self.client.get('/equipos/',{'todos':'1'}).context['page'].paginator.count,2)
    def test_solicitudes_compartidas(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.other,tipo='Correctivo',descripcion='Solicitud común')
        self.client.force_login(self.user)
        response=self.client.get('/solicitudes/')
        self.assertContains(response,sol.codigo)
        self.assertContains(response,'Realizar mantenimiento')
        self.assertContains(response,f'data-row-url="/solicitudes/{sol.pk}/"')
        self.assertEqual(self.client.get(f'/solicitudes/{sol.pk}/').status_code,200)

    def test_cualquier_usuario_toma_y_finaliza_sin_asignacion_previa(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.other,tipo='MIC',descripcion='Falla')
        self.client.force_login(self.user)
        url=f'/solicitudes/{sol.pk}/realizar/'
        self.assertEqual(self.client.get(url).status_code,405)
        self.assertRedirects(self.client.post(url),f'/mantenimientos/nuevo/?solicitud={sol.pk}')
        sol.refresh_from_db()
        self.assertEqual(sol.tecnico,self.user)
        self.assertEqual(sol.estado,'Asignada')
        self.client.post(url)
        sol.refresh_from_db(); self.assertEqual(sol.version,2)
        self.client.force_login(self.other)
        self.client.post(url)
        sol.refresh_from_db(); self.assertEqual(sol.tecnico,self.user)
        self.client.force_login(self.user)
        response=self.client.post('/mantenimientos/nuevo/',{'equipo':self.equipo.pk,'solicitud':sol.pk,'fecha':'2026-10-13','tipo':'MIC','diagnostico':'Revisión del equipo'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.client.get(response.url).status_code,200)
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Cerrada')
        self.client.post(url)
        self.assertEqual(Mantenimiento.objects.filter(solicitud=sol).count(),1)
    def test_edicion_obsoleta(self):
        self.equipo.laboratorio.responsable=self.user
        self.equipo.laboratorio.save(update_fields=['responsable'])
        self.grant('change_equipo'); self.client.force_login(self.user)
        response=self.client.post(f'/equipos/{self.equipo.pk}/editar/',{'descripcion':'Cambio','laboratorio':self.equipo.laboratorio_id,'version':0})
        self.assertContains(response,'Otro usuario')
        self.equipo.refresh_from_db(); self.assertEqual(self.equipo.descripcion,'Multímetro')

    def test_tomar_solicitud_historica_sin_equipo(self):
        sol=Solicitud.objects.create(solicitante=self.other,tipo='MIC',descripcion='Histórica',datos_legacy={'equipo':'Equipo original'})
        self.client.force_login(self.user)
        response=self.client.get(f'/solicitudes/{sol.pk}/')
        self.assertContains(response,'Realizar mantenimiento')
        self.assertNotContains(response,'Primero debe vincularse')
        self.assertNotContains(response,'Técnico: Sin asignar')
        self.client.post(f'/solicitudes/{sol.pk}/realizar/')
        sol.refresh_from_db(); self.assertEqual(sol.tecnico,self.user)
        self.assertIsNone(sol.equipo_id)
        response=self.client.post('/mantenimientos/nuevo/',{'equipo':self.equipo.pk,'solicitud':sol.pk,'fecha':'2026-10-13','tipo':'MIC','diagnostico':'Revisión del equipo'})
        self.assertEqual(response.status_code,302)
        sol.refresh_from_db()
        self.assertEqual(sol.estado,'Cerrada')
        self.assertIsNone(sol.equipo_id)
        self.assertEqual(sol.datos_legacy,{'equipo':'Equipo original'})

    def test_formulario_rapido_reutiliza_datos_historicos(self):
        from django.utils import timezone
        sol=Solicitud.objects.create(solicitante=self.other,tecnico=self.user,estado='Asignada',tipo='MEPV',descripcion='Revisar',datos_legacy={'equipo':'Equipo original','placa':'123'})
        self.client.force_login(self.user)
        url=f'/mantenimientos/nuevo/?solicitud={sol.pk}'
        response=self.client.get(url)
        self.assertContains(response,'Registrar trabajo realizado')
        self.assertContains(response,'Equipo original')
        self.assertContains(response,'123')
        self.assertNotContains(response,'Vincular equipo')
        self.assertTrue(response.context['form']['equipo'].is_hidden)
        self.assertContains(response,timezone.localdate().isoformat())
        self.assertTrue(response.context['form'].fields['solicitud'].disabled)
        self.assertTrue(response.context['form'].fields['tipo'].disabled)
        response=self.client.post(url,{'especificaciones':'Limpieza y revisión','estado':'Operativo','tipo':'MIC','diagnostico':'Revisión del equipo'})
        self.assertEqual(response.status_code,302)
        m=Mantenimiento.objects.get(solicitud=sol)
        self.assertEqual(m.tipo,'MEPV')
        self.assertEqual(m.fecha,timezone.localdate())
        self.assertIsNone(m.equipo_id)
        self.assertEqual(m.datos_legacy['equipo'],'Equipo original')
        self.assertEqual(self.client.get(response.url).status_code,200)
    def test_cierre_solicitud(self):
        self.grant('add_mantenimiento'); self.client.force_login(self.user)
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.other,tecnico=self.user,tipo='Correctivo',descripcion='Falla',estado='Asignada')
        payload={'equipo':self.equipo.pk,'solicitud':sol.pk,'fecha':'2026-10-02','tipo':'Correctivo','diagnostico':'Revisión del equipo'}
        self.assertEqual(self.client.post('/mantenimientos/nuevo/',payload).status_code,302)
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Cerrada')
        self.client.post('/mantenimientos/nuevo/',payload)
        self.assertEqual(Mantenimiento.objects.count(),1)
    def test_documento_excel(self):
        from .documents import generar
        from openpyxl import load_workbook
        wb=load_workbook(generar(self.equipo,'xlsx'))
        self.assertEqual(wb.active['H13'].value,'001'); wb.close()

    def test_importacion_simulada_e_idempotente(self):
        from pathlib import Path
        from tempfile import TemporaryDirectory
        from openpyxl import Workbook
        from django.core.management import call_command
        with TemporaryDirectory() as temp:
            root=Path(temp); resources=root/'resources'; resources.mkdir()
            inventories=resources/'inventarios'; inventories.mkdir()
            wb=Workbook(); ws=wb.active; ws.title='GIF-PR-003-FR-004'
            for _ in range(14): ws.append([None])
            ws.append(['Placa UD','Serial','Descripci\u00f3n','N. Interno'])
            ws.append(['0001','S-1','Equipo de prueba','A1'])
            wb.save(inventories/'603-prueba.xlsx'); wb.close()
            with override_settings(RESOURCES_ROOT=resources):
                options={'user':self.user.username,'report':str(root/'report.json')}
                call_command('importar_legacy',**options)
                self.assertFalse(Equipo.objects.filter(placa='0001').exists())
                call_command('importar_legacy',apply=True,**options)
                call_command('importar_legacy',apply=True,**options)
                self.assertEqual(Equipo.objects.filter(placa='0001').count(),1)


class MigracionCompletaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin=User.objects.create(username='admin',is_superuser=True,is_staff=True)
        cls.lector=User.objects.create(username='lector')
        cls.tecnico=User.objects.create(username='tecnico')
        cls.lab=Laboratorio.objects.create(nombre='603 - Electrónica A')
        cls.equipo=Equipo.objects.create(placa='0007',numero_interno='EA07',serial='ABC',descripcion='Osciloscopio',laboratorio=cls.lab)
        cls.lector.user_permissions.add(*Permission.objects.filter(content_type__app_label='lab',codename__in=['view_equipo','view_mantenimiento','add_solicitud']))
        cls.tecnico.user_permissions.add(*Permission.objects.filter(content_type__app_label='lab',codename__in=['view_equipo','add_mantenimiento','view_mantenimiento']))

    def test_pantallas_del_sistema(self):
        self.client.force_login(self.admin)
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.lector,tipo='MIC',descripcion='Falla')
        mt=Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.tecnico,fecha=date(2026,10,2),tipo='MIPV')
        urls=['/','/equipos/','/hojas-vida/','/equipos/nuevo/',f'/equipos/{self.equipo.pk}/',f'/equipos/{self.equipo.pk}/editar/',f'/equipos/{self.equipo.pk}/hoja-vida/',f'/equipos/{self.equipo.pk}/fotografia/','/solicitudes/',f'/solicitudes/{sol.pk}/','/solicitudes/nueva/','/mantenimientos/','/mantenimientos/nuevo/',f'/mantenimientos/{mt.pk}/',f'/mantenimientos/{mt.pk}/editar/','/mantenimientos/estadisticas/','/laboratorios/','/laboratorios/nuevo/','/importacion/','/importacion/registros/','/documentos/','/perfil/','/usuarios/','/usuarios/nuevo/',f'/usuarios/{self.lector.pk}/editar/',f'/usuarios/{self.lector.pk}/password/','/auditoria/','/accounts/password_change/']
        for url in urls:
            with self.subTest(url=url): self.assertEqual(self.client.get(url).status_code,200)

    def test_cancelacion_y_privacidad(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.lector,tipo='MIC',descripcion='Falla')
        self.client.force_login(self.tecnico)
        self.assertEqual(self.client.post(f'/solicitudes/{sol.pk}/cancelar/',{'version':1,'motivo':'Ajeno'}).status_code,403)
        self.client.force_login(self.lector)
        self.assertEqual(self.client.post(f'/solicitudes/{sol.pk}/cancelar/',{'version':1,'motivo':'Ya no se requiere'}).status_code,302)
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Cancelada')

    def test_asignacion_version_y_cierre(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.lector,tipo='MIC',descripcion='Falla')
        self.client.force_login(self.admin)
        self.client.post(f'/solicitudes/{sol.pk}/asignar/',{'version':1,'tecnico':self.tecnico.pk})
        sol.refresh_from_db(); self.assertIsNone(sol.tecnico)
        self.client.force_login(self.tecnico)
        self.client.post(f'/solicitudes/{sol.pk}/realizar/')
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Asignada'); self.assertEqual(sol.version,2)
        self.client.post(f'/solicitudes/{sol.pk}/cancelar/',{'version':1,'motivo':'Formulario viejo'})
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Asignada')
        self.client.force_login(self.tecnico)
        payload={'solicitud':sol.pk,'equipo':self.equipo.pk,'tipo':'MIC','diagnostico':'Revisión del equipo','fecha':'2026-10-02','estado':'Operativo'}
        self.client.post('/mantenimientos/nuevo/',payload); self.client.post('/mantenimientos/nuevo/',payload)
        self.assertEqual(Mantenimiento.objects.filter(solicitud=sol).count(),1)
        sol.refresh_from_db(); self.assertEqual(sol.estado,'Cerrada')

    def test_tecnico_no_atiende_solicitud_ajena(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.lector,tecnico=self.admin,tipo='MIC',descripcion='Falla',estado='Asignada')
        self.client.force_login(self.tecnico)
        self.client.post('/mantenimientos/nuevo/',{'solicitud':sol.pk,'equipo':self.equipo.pk,'tipo':'MIC','diagnostico':'Revisión del equipo','fecha':'2026-10-02'})
        self.assertFalse(Mantenimiento.objects.exists())

    def test_solicitud_con_datos_del_escritorio(self):
        self.client.force_login(self.lector)
        response=self.client.post('/solicitudes/nueva/',{'equipo':self.equipo.pk,'tipo':'MEPV','nombre_solicitante':'Estudiante prueba','documento':'00123','celular':'3000000000','tipo_solicitante':'Estudiante','fecha_reporte':'2026-10-02','hora_reporte':'08:30','descripcion':'Revisión preventiva','asignatura':'Circuitos','banco':'2'})
        self.assertEqual(response.status_code,302)
        obj=Solicitud.objects.get(); self.assertEqual(obj.documento,'00123'); self.assertEqual(obj.tipo,'MEPV')

    def test_identidades_en_conflicto_no_se_asocian(self):
        from .importing import matching_equipment
        eq=Equipo.objects.create(placa='OTRA',numero_interno='OTRO',descripcion='Otro',laboratorio=self.lab)
        self.assertIsNone(matching_equipment('0007','OTRO'))
        self.assertEqual(matching_equipment('0007','EA07').pk,self.equipo.pk)

    def test_importacion_encabezado_tres_finalizada_y_sin_vinculo(self):
        from pathlib import Path
        from tempfile import TemporaryDirectory
        from openpyxl import Workbook
        from .importing import importar
        from .models import RegistroLegacy
        with TemporaryDirectory() as directory:
            root=Path(directory)/'resources'; folder=root/'mantenimientos'; folder.mkdir(parents=True)
            wb=Workbook(); ws=wb.active; ws.title='Datos'; ws.append(['Técnico']); ws.append([1]); ws.append(['CONSECUTIVO']+[f'Col {i}' for i in range(2,28)])
            for placa in ['0007','NO-EXISTE']:
                row=[None]*27; row[0]='S-'+placa; row[1]='Persona histórica'; row[3]='MIP'; row[4]=date(2024,2,3); row[6]='Equipo histórico'; row[8]=placa; row[13]='Trabajo'; row[15]='Técnico histórico'; row[16]=date(2024,2,4); row[20]='Intervención'; row[24]='Finalizada'; ws.append(row)
            wb.save(folder/'REPORTE_MANTENIMIENTOS.xlsx'); wb.close()
            with override_settings(RESOURCES_ROOT=root):
                report=importar(self.admin,apply=True); importar(self.admin,apply=True)
            self.assertEqual(Solicitud.objects.count(),2); self.assertEqual(Mantenimiento.objects.count(),2)
            self.assertEqual(Solicitud.objects.filter(estado='Cerrada').count(),2)
            self.assertEqual(Solicitud.objects.filter(equipo__isnull=True).count(),1)
            self.assertEqual(RegistroLegacy.objects.count(),2); self.assertEqual(len(report['incidencias']),1)

    def test_conciliar_historia_actualiza_ambos_registros(self):
        from .models import RegistroLegacy
        record=RegistroLegacy.objects.create(origen='reporte:4',archivo='reporte.xlsx',hoja='Datos',fila=4,categoria='reporte')
        sol=Solicitud.objects.create(solicitante=self.admin,tipo='MIC',descripcion='Histórico',origen=record.origen)
        mt=Mantenimiento.objects.create(tecnico=self.admin,tipo='MIC',origen=record.origen)
        self.client.force_login(self.admin)
        response=self.client.post(f'/importacion/registros/{record.pk}/',{'equipo':self.equipo.pk,'nota':'Verificado físicamente'})
        self.assertEqual(response.status_code,302)
        sol.refresh_from_db(); mt.refresh_from_db(); record.refresh_from_db()
        self.assertEqual(sol.equipo_id,self.equipo.pk); self.assertEqual(mt.equipo_id,self.equipo.pk); self.assertEqual(record.estado,'Revisada')

    def test_pdf_paginado_con_todos_los_mantenimientos(self):
        from .documents import generar
        from pypdf import PdfReader
        for i in range(12): Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.tecnico,fecha=date(2026,10,2),tipo='MIC',especificaciones=f'TRABAJO-{i:02d} '+('Se revisaron los componentes. '*15))
        pdf=PdfReader(generar(self.equipo,'pdf'))
        content='\n'.join(page.extract_text() for page in pdf.pages)
        self.assertGreater(len(pdf.pages),1)
        for i in range(12): self.assertIn(f'TRABAJO-{i:02d}',content)
        self.assertIn('0007',content)

    def test_excel_mas_de_siete_intervenciones(self):
        from .documents import generar
        from openpyxl import load_workbook
        for i in range(10): Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.tecnico,fecha=date(2026,10,2),tipo='MIC',especificaciones=f'Trabajo {i}')
        wb=load_workbook(generar(self.equipo,'xlsx')); ws=wb.active
        self.assertEqual(ws['O39'].value,'Trabajo 9'); wb.close()

    def test_foto_privada_y_subida(self):
        from io import BytesIO
        from tempfile import TemporaryDirectory
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        image=BytesIO(); Image.new('RGB',(30,30),'blue').save(image,format='PNG')
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.client.force_login(self.admin)
            response=self.client.post(f'/equipos/{self.equipo.pk}/fotografia/',{'foto':SimpleUploadedFile('photo.png',image.getvalue(),content_type='image/png')})
            self.assertEqual(response.status_code,302)
            response=self.client.get(f'/equipos/{self.equipo.pk}/foto/'); self.assertEqual(response.status_code,200); [closer() for closer in list(response._resource_closers)]
            self.client.logout(); self.assertEqual(self.client.get(f'/equipos/{self.equipo.pk}/foto/').status_code,302)

    def test_lote_y_descarga_solo_del_propietario(self):
        from tempfile import TemporaryDirectory
        from .jobs import process
        from .models import TrabajoDocumento
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            obj=TrabajoDocumento.objects.create(usuario=self.lector,formato='pdf',equipos=[self.equipo.pk])
            process(obj.pk); process(obj.pk); obj.refresh_from_db()
            self.assertEqual(obj.estado,'Completado')
            self.client.force_login(self.tecnico); self.assertEqual(self.client.get(f'/documentos/{obj.pk}/descargar/').status_code,404)
            self.client.force_login(self.lector); response=self.client.get(f'/documentos/{obj.pk}/descargar/'); self.assertEqual(response.status_code,200); [closer() for closer in list(response._resource_closers)]

    def test_reporte_excel_preserva_columnas(self):
        from io import BytesIO
        from openpyxl import load_workbook
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.lector,tipo='MIP',descripcion='Falla',documento='00123',estado='Cerrada')
        Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.tecnico,tipo='MIPV',fecha=date(2026,10,2),especificaciones='Preventivo independiente')
        self.client.force_login(self.admin)
        response=self.client.get('/solicitudes/reporte/')
        wb=load_workbook(BytesIO(b''.join(response.streaming_content))); ws=wb['Datos']
        self.assertEqual(ws['C4'].value,'00123'); self.assertEqual(ws['Y4'].value,'Finalizada'); self.assertEqual(ws['U5'].value,'Preventivo independiente'); self.assertIn('CPS',wb.sheetnames); wb.close()

    def test_exportacion_formula_se_guarda_como_texto(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.equipo.descripcion='=HYPERLINK("https://example.com")'; self.equipo.save()
        self.client.force_login(self.lector); response=self.client.get('/equipos/exportar/')
        wb=load_workbook(BytesIO(b''.join(response.streaming_content))); self.assertEqual(wb.active['D2'].data_type,'s'); wb.close()

    def test_csrf_en_operaciones(self):
        from django.test import Client
        client=Client(enforce_csrf_checks=True); client.force_login(self.admin)
        self.assertEqual(client.post('/solicitudes/nueva/',{}).status_code,403)

    def test_usuario_sin_rol_no_modifica_ni_importa(self):
        self.client.force_login(self.lector)
        for url in ['/equipos/nuevo/','/importacion/','/usuarios/nuevo/',f'/equipos/{self.equipo.pk}/editar/']:
            self.assertEqual(self.client.get(url).status_code,403)

    def test_lector_de_historial_no_ve_contactos_originales(self):
        obj=Mantenimiento.objects.create(equipo=self.equipo,tecnico=self.tecnico,tipo='MIC',datos_legacy={'documento':'DOCUMENTO-PRIVADO-123'})
        self.client.force_login(self.lector)
        self.assertNotContains(self.client.get(f'/mantenimientos/{obj.pk}/'),'DOCUMENTO-PRIVADO-123')

    def test_no_administrador_no_cambia_superusuario(self):
        self.tecnico.user_permissions.add(Permission.objects.get(content_type__app_label='auth',codename='change_user'))
        self.client.force_login(self.tecnico)
        self.assertEqual(self.client.get(f'/usuarios/{self.admin.pk}/editar/').status_code,403)


@skipUnless(connection.vendor=='postgresql','La concurrencia real se verifica con PostgreSQL.')
class ConcurrenciaPostgreSQLTests(TransactionTestCase):
    def setUp(self):
        self.user=User.objects.create(username='tecnico',is_superuser=True)
        self.equipo=Equipo.objects.create(descripcion='Equipo concurrente',laboratorio=Laboratorio.objects.create(nombre='603'))
    def race(self,url,payloads):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from django.test import Client
        barrier=Barrier(2)
        def worker(payload):
            close_old_connections()
            try:
                client=Client();client.force_login(User.objects.get(pk=self.user.pk));barrier.wait(timeout=10)
                return client.post(url,payload).status_code
            finally: close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as executor: return list(executor.map(worker,payloads))
    def test_dos_tecnicos_no_cierran_dos_veces(self):
        sol=Solicitud.objects.create(equipo=self.equipo,solicitante=self.user,tecnico=self.user,tipo='MIC',descripcion='Falla',estado='Asignada')
        payload={'equipo':self.equipo.pk,'solicitud':sol.pk,'fecha':'2026-10-02','tipo':'MIC','diagnostico':'Revisión del equipo','estado':'Operativo'}
        result=self.race('/mantenimientos/nuevo/',[payload,payload])
        self.assertEqual(sorted(result),[200,302]);self.assertEqual(Mantenimiento.objects.filter(solicitud=sol).count(),1)
    def test_dos_ediciones_no_sobrescriben_la_primera(self):
        payload={'laboratorio':self.equipo.laboratorio_id,'descripcion':'Cambio uno','version':1}
        other={**payload,'descripcion':'Cambio dos'}
        result=self.race(f'/equipos/{self.equipo.pk}/editar/',[payload,other])
        self.assertEqual(sorted(result),[200,302]);self.equipo.refresh_from_db();self.assertEqual(self.equipo.version,2)
